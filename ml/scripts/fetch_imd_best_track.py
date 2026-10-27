"""Phase 2 / Step 2: get the official IMD RSMC best-track workbook and verify it.

What it does
  1. Downloads the workbook from the RSMC New Delhi site (or registers a file you
     downloaded by hand with --inspect-only).
  2. Checks it is a real .xlsx (not an HTML error page).
  3. Writes SOURCE.txt next to it (URL, date, size, SHA-256) for provenance.
  4. Prints every sheet's size and the first rows, so we can see the real layout
     before writing a parser (next step). It does NOT modify the raw file.

Usage (repo root, Windows PowerShell or Colab):
    pip install openpyxl
    python ml/scripts/fetch_imd_best_track.py
    python ml/scripts/fetch_imd_best_track.py --inspect-only     # file downloaded manually
    python ml/scripts/fetch_imd_best_track.py --root /content/drive/MyDrive/trinetra_data
"""
import argparse
import hashlib
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path

# Link taken from https://rsmcnewdelhi.imd.gov.in/report.php?internal_menu=MzM=
# The file name contains a hash and can change when IMD uploads a new version.
# If this URL fails, download by hand from that page and use --inspect-only.
DEFAULT_URL = (
    "https://rsmcnewdelhi.imd.gov.in/download.php"
    "?path=uploads/best-track/78b4b0_Best_Tracks__Data__1982-2026_.xlsx"
)
OUT_NAME = "imd_best_tracks.xlsx"


def download(url: str, dest: Path) -> None:
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (TRINETRA data fetch)"})
    tmp = dest.with_suffix(".part")
    with urllib.request.urlopen(req, timeout=120) as resp, open(tmp, "wb") as f:
        while chunk := resp.read(1 << 20):
            f.write(chunk)
    tmp.replace(dest)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def inspect(path: Path, n_rows: int = 6) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"\nSheets in {path.name}: {wb.sheetnames}")
    for ws in wb.worksheets:
        print(f"\n=== sheet '{ws.title}'  rows={ws.max_row}  cols={ws.max_column}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= n_rows:
                break
            print(i + 1, [c for c in row][:12])


def main() -> None:
    default_root = Path(__file__).resolve().parents[1] / "data"
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", type=Path, default=default_root)
    ap.add_argument("--url", default=DEFAULT_URL)
    ap.add_argument("--inspect-only", action="store_true",
                    help="skip download; use the .xlsx already in raw/imd_best_track/")
    ap.add_argument("--force", action="store_true", help="re-download even if present")
    args = ap.parse_args()

    folder = args.root / "raw" / "imd_best_track"
    folder.mkdir(parents=True, exist_ok=True)
    dest = folder / OUT_NAME

    if args.inspect_only:
        found = [p for p in folder.glob("*.xlsx")]
        if len(found) != 1:
            raise SystemExit(f"Expected exactly one .xlsx in {folder}, found {len(found)}.")
        dest = found[0]
        source = "manual download from https://rsmcnewdelhi.imd.gov.in/report.php?internal_menu=MzM="
    else:
        if dest.exists() and not args.force:
            print(f"{dest} already exists (use --force to re-download).")
        else:
            print(f"Downloading {args.url}")
            download(args.url, dest)
        source = args.url

    if not zipfile.is_zipfile(dest):
        raise SystemExit(f"{dest} is not a valid .xlsx (probably an HTML error page). "
                         "Download it manually and rerun with --inspect-only.")

    (folder / "SOURCE.txt").write_text(
        f"file: {dest.name}\nsource: {source}\n"
        f"retrieved_utc: {datetime.now(timezone.utc).isoformat(timespec='seconds')}\n"
        f"size_bytes: {dest.stat().st_size}\nsha256: {sha256(dest)}\n",
        encoding="utf-8",
    )
    print(f"OK  {dest}  ({dest.stat().st_size/1e6:.2f} MB); wrote SOURCE.txt")
    inspect(dest)


if __name__ == "__main__":
    main()
